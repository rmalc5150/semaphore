#!/usr/bin/env python3
"""Collect public UK MHRA ILAP and orphan-medicine records by day/month.

Official spreadsheets are parsed in memory and JSON is streamed directly to S3.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import subprocess
import sys
from datetime import date, datetime
from html.parser import HTMLParser
from typing import Any
from urllib.parse import urljoin
from urllib.request import Request, urlopen

from decodo_scraping_api import scrape_url

from clinicaltrials_studies import parse_date, period_bounds

ILAP_PAGE = "https://www.gov.uk/government/publications/innovative-licensing-and-access-pathway-ilap"
ORPHAN_PAGE = "https://www.gov.uk/government/publications/orphan-registered-medicinal-products"
S3_PREFIX = "s3://processed-059416028963-us-east-1-an/regulatory-designations/uk/"
USER_AGENT = "mhra-designations/1.0"


class Links(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[tuple[str, str]] = []
        self.href: str | None = None
        self.text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self.href = dict(attrs).get("href")
            self.text = []

    def handle_data(self, data: str) -> None:
        if self.href:
            self.text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self.href:
            self.links.append((self.href, " ".join(self.text).strip()))
            self.href = None


def download(url: str) -> tuple[bytes, str]:
    request = Request(url, headers={"User-Agent": USER_AGENT, "Accept": "*/*"})
    with urlopen(request, timeout=90) as response:
        return response.read(), response.geturl()


def attachment(page: str, suffix: str, text_contains: str) -> str:
    raw, final_url = scrape_url(page, proxy_pool="standard").encode(), page
    parser = Links()
    parser.feed(raw.decode("utf-8", errors="replace"))
    matches = [urljoin(final_url, href) for href, text in parser.links
               if href.lower().split("?")[0].endswith(suffix) and text_contains.lower() in text.lower()]
    if not matches:
        raise RuntimeError(f"could not find {text_contains!r} {suffix} attachment on {page}")
    return matches[-1]


def clean(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value.strip() if isinstance(value, str) else value


def record_date(record: dict[str, Any]) -> date | None:
    for key, value in record.items():
        if "expir" in key.lower():
            continue
        if not any(word in key.lower() for word in ("date", "authoris", "award")):
            continue
        if isinstance(value, (date, datetime)):
            return value.date() if isinstance(value, datetime) else value
        if isinstance(value, str):
            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d %B %Y", "%d-%b-%Y"):
                try:
                    return datetime.strptime(value.strip(), fmt).date()
                except ValueError:
                    pass
    return None


def previous_snapshot() -> dict[str, Any] | None:
    source = S3_PREFIX + "snapshots/current.json"
    result = subprocess.run(["aws", "s3", "cp", source, "-", "--only-show-errors"],
                            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True)
    if result.returncode:
        return None
    try:
        value = json.loads(result.stdout)
        return value if isinstance(value, dict) else None
    except json.JSONDecodeError:
        return None


def identity(category: str, record: dict[str, Any]) -> str:
    preferred = ("Designation number", "Product Name", "Product name", "Company Name*")
    parts = [str(record.get(key, "")).strip().lower() for key in preferred if record.get(key)]
    return category + ":" + ("|".join(parts) or json.dumps(record, sort_keys=True, ensure_ascii=False))


def observed_changes(previous: dict[str, Any] | None, current: dict[str, list[dict[str, Any]]],
                     observed: date) -> list[dict[str, Any]]:
    if previous is None:
        return []  # Establish a baseline rather than reporting every existing row as new.
    events = []
    for category, native, key in (("BREAKTHROUGH", "ILAP Innovation Passport", "ilap"),
                                  ("ORPHAN", "UK orphan marketing authorisation", "orphanRegister")):
        old = {identity(category, r): r for r in previous.get(key, []) if isinstance(r, dict)}
        new = {identity(category, r): r for r in current[key]}
        for event_type, keys in (("OBSERVED_ADDED", new.keys() - old.keys()),
                                 ("OBSERVED_REMOVED", old.keys() - new.keys()),
                                 ("OBSERVED_CHANGED", new.keys() & old.keys())):
            for item_id in keys:
                if event_type == "OBSERVED_CHANGED" and new[item_id] == old[item_id]:
                    continue
                events.append({"eventDate": observed.isoformat(), "regulator": "MHRA", "jurisdiction": "UK",
                               "normalizedCategory": category, "nativeProgram": native,
                               "eventType": event_type, "record": new.get(item_id) or old[item_id],
                               "previousRecord": old.get(item_id) if event_type == "OBSERVED_CHANGED" else None})
    return events


def read_ilap() -> tuple[list[dict[str, Any]], str]:
    from openpyxl import load_workbook

    url = attachment(ILAP_PAGE, ".xlsx", "Innovation Passport")
    raw, _ = download(url)
    sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    header_at = next((i for i, row in enumerate(rows) if sum(v is not None for v in row) >= 2), None)
    if header_at is None:
        return [], url
    headers = [str(v).strip() if v is not None else f"column_{i+1}" for i, v in enumerate(rows[header_at])]
    records = []
    for row in rows[header_at + 1:]:
        if not any(v not in (None, "") for v in row):
            continue
        records.append({headers[i]: clean(v) for i, v in enumerate(row) if i < len(headers) and v not in (None, "")})
    return records, url


def read_orphans() -> tuple[list[dict[str, Any]], str]:
    url = attachment(ORPHAN_PAGE, ".csv", "Orphan Register")
    raw, _ = download(url)
    text = raw.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    header = next((i for i, line in enumerate(lines) if line.lstrip().startswith("Product name ,")), None)
    if header is None:
        raise RuntimeError("could not find the MHRA Orphan Register CSV header")
    records = [{str(k).strip(): clean(v) for k, v in row.items() if k and v not in (None, "")}
               for row in csv.DictReader(io.StringIO("\n".join(lines[header:])))]
    return records, url


def upload(payload: dict[str, Any], key: str) -> str:
    destination = S3_PREFIX + key.lstrip("/")
    process = subprocess.Popen(["aws", "s3", "cp", "-", destination, "--only-show-errors"],
                               stdin=subprocess.PIPE, text=True, encoding="utf-8")
    assert process.stdin is not None
    json.dump(payload, process.stdin, ensure_ascii=False, indent=2)
    process.stdin.write("\n"); process.stdin.close()
    if process.wait():
        raise RuntimeError("AWS CLI upload failed")
    return destination


def main() -> int:
    if len(sys.argv) > 1 and sys.argv[1] in {"historical", "daily"}:
        from regulatory_collection import main as shared_main
        return shared_main('uk', sys.argv[1:])
    parser = argparse.ArgumentParser(description="Collect UK MHRA ILAP and orphan designations.")
    parser.add_argument("period", choices=("day", "month"))
    parser.add_argument("date", type=parse_date, help="YYYY-MM-DD (or YYYY-MM for month)")
    parser.add_argument("--s3-key")
    args = parser.parse_args()
    start, end = period_bounds(args.period, args.date)
    try:
        ilap, ilap_url = read_ilap()
        orphans, orphan_url = read_orphans()
        current = {"ilap": ilap, "orphanRegister": orphans}
        events = []
        for category, native, records in (("BREAKTHROUGH", "ILAP Innovation Passport", ilap),
                                           ("ORPHAN", "UK orphan marketing authorisation", orphans)):
            for record in records:
                event_date = record_date(record)
                if event_date and start <= event_date <= end:
                    events.append({"eventDate": event_date.isoformat(), "regulator": "MHRA",
                                   "jurisdiction": "UK", "normalizedCategory": category,
                                   "nativeProgram": native, "eventType": "DESIGNATED", "record": record})
        observed = date.today()
        if start <= observed <= end:
            events.extend(observed_changes(previous_snapshot(), current, observed))
        payload = {"regulator": "MHRA", "period": args.period, "startDate": str(start), "endDate": str(end),
                   "sources": {"ilap": ilap_url, "orphanRegister": orphan_url},
                   "coverageNote": "The UK grants orphan status at marketing authorisation; ILAP is a close counterpart, not an exact FDA Breakthrough Therapy equivalent.",
                   "eventCount": len(events), "events": events,
                   "snapshots": {"ilap": ilap, "orphanRegister": orphans}}
        key = args.s3_key or f"mhra_designations_{args.period}_{start}_{end}.json"
        destination = upload(payload, key)
        upload(current, "snapshots/current.json")
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr); return 1
    print(f"Uploaded {len(events)} UK designation events to {destination}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
