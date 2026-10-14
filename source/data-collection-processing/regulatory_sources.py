"""Source-specific pagination, identities and date semantics for group one."""
from datetime import date, datetime, timedelta
import csv
import io
import os
import re
from html import unescape
from urllib.parse import urlencode, urljoin

from regulatory_collection import SourceIssue, require
from research_contract import digest


def day(value):
    if not value:
        return None
    for fmt, text in (("%Y-%m-%d", str(value)[:10]), ("%d/%m/%Y", str(value)),
                      ("%m/%d/%Y", str(value)), ("%Y%m%d", str(value))):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def in_window(value, start, end):
    return value is not None and start.isoformat() <= value < end.isoformat()


def item(identity, kind, data, url, basis, event=None, names=(), title=None, **extra):
    require(bool(identity), "missing_provider_identity")
    return {"id": str(identity), "kind": kind, "data": data, "url": url, "date_basis": basis,
            "date": event, "company_names": [x for x in names if isinstance(x, str) and x],
            "title": title, **extra}


def validate_total(response, rows, state, field="totalCount"):
    total = response.get(field)
    require(isinstance(total, int) and not isinstance(total, bool) and total >= 0, "missing_source_total")
    if "total" in state:
        require(total == state["total"], "source_changed_during_pagination")
    count = state.get("count", 0) + len(rows)
    require(count <= total and (rows or count == total), "source_count_mismatch")
    return total, count


class ClinicalTrials:
    source = "clinicaltrials_gov"
    date_basis = "last_update_post_date"
    coverage = "Current study versions whose last-update date matches; not historical snapshots. Global registry, all sponsor sectors."

    def page(self, client, start, end, mode, state, size):
        params = {"format": "json", "query.term": f"AREA[LastUpdatePostDate]RANGE[{start}, {end-timedelta(days=1)}]",
                  "pageSize": str(size), "countTotal": "true", "sort": "LastUpdatePostDate:asc"}
        if state.get("token"):
            params["pageToken"] = state["token"]
        response = client.json("https://clinicaltrials.gov/api/v2/studies", params=params)
        rows = response.get("studies")
        require(isinstance(rows, list), "missing_studies")
        # The API returns totalCount on the first page only, even with
        # countTotal=true on subsequent token requests.
        if "totalCount" not in response and state.get("token") and "total" in state:
            response = {**response, "totalCount": state["total"]}
        total, count = validate_total(response, rows, state)
        token = response.get("nextPageToken")
        require(not token or token != state.get("token"), "pagination_cycle")
        require(bool(token) == (count < total), "source_count_mismatch")
        output = []
        for row in rows:
            protocol = row["protocolSection"]
            identity = protocol["identificationModule"]["nctId"]
            require(re.fullmatch(r"NCT\d{8}", identity), "invalid_trial_identifier")
            updated = day(protocol["statusModule"].get("lastUpdatePostDateStruct", {}).get("date"))
            require(in_window(updated, start, end), "study_outside_update_window")
            sponsors = protocol.get("sponsorCollaboratorsModule", {})
            names = [sponsors.get("leadSponsor", {}).get("name"),
                     protocol["identificationModule"].get("organization", {}).get("fullName")]
            names += [x.get("name") for x in sponsors.get("collaborators", [])]
            output.append(item(identity, "clinical_trial", row, f"https://clinicaltrials.gov/study/{identity}",
                self.date_basis, updated, names, protocol["identificationModule"].get("briefTitle"),
                updated=updated, coverage_note=self.coverage))
        return output, {"token": token, "total": total, "count": count}, not token


class FDA:
    source = "fda_openfda"
    date_basis = "submission_action_or_letter_date"
    coverage = "Public approvals and CRLs; letter/action dates do not establish publication. CRLs include approved and unapproved applications."

    def page(self, client, start, end, mode, state, size):
        stage, offset = state.get("stage", "approvals"), state.get("count", 0)
        if offset > 25000:
            raise SourceIssue("openfda_pagination_limit", True)
        url = "https://api.fda.gov/drug/drugsfda.json" if stage == "approvals" else "https://api.fda.gov/transparency/crl.json"
        params = {"limit": str(min(size, 1000)), "skip": str(offset)}
        if stage == "approvals":
            params["search"] = f"submissions.submission_status_date:[{start:%Y%m%d} TO {end-timedelta(days=1):%Y%m%d}]"
            params["sort"] = "application_number:asc"
        response = client.json(url, params=params, api_key=os.getenv("OPENFDA_API_KEY"), no_matches=True)
        rows = response.get("results")
        require(isinstance(rows, list), "missing_fda_results")
        total, count = validate_total({"totalCount": response.get("meta", {}).get("results", {}).get("total")}, rows, state)
        revision = response.get("meta", {}).get("last_updated")
        if state.get("revision"):
            require(revision == state["revision"], "source_changed_during_pagination")
        output = []
        for row in rows:
            if stage == "approvals":
                require(row.get("application_number"), "missing_application_number")
                for submission in row.get("submissions", []):
                    if submission.get("submission_status") != "AP":
                        continue
                    event = day(submission.get("submission_status_date"))
                    require(event, "invalid_fda_action_date")
                    if not in_window(event, start, end):
                        continue
                    require(submission.get("submission_type") and submission.get("submission_number"), "missing_submission_identity")
                    identity = "approval:" + ":".join(str(x) for x in (row["application_number"], submission["submission_type"], submission["submission_number"]))
                    output.append(item(identity, "regulatory_event", {"action": "APPROVAL", "application": row,
                        "submission": submission}, url, self.date_basis, event, [row.get("sponsor_name")],
                        row["application_number"], identity_basis="documented_composite", coverage_note=self.coverage))
            else:
                if str(row.get("letter_type") or "").upper() != "COMPLETE RESPONSE":
                    continue
                try:
                    event = datetime.strptime(row["letter_date"], "%m/%d/%Y").date().isoformat()
                except (ValueError, KeyError):
                    raise SourceIssue("invalid_crl_date") from None
                if in_window(event, start, end):
                    require(row.get("file_name"), "missing_crl_identity")
                    output.append(item("crl:" + row["file_name"], "regulatory_event", row, url,
                        self.date_basis, event, [row.get("company_name")], row["file_name"], coverage_note=self.coverage))
        exhausted = count == total
        following = {"stage": stage, "count": count, "total": total, "revision": revision}
        if exhausted and stage == "approvals":
            return output, {"stage": "crls"}, False
        return output, following, exhausted


class EMA:
    source = "ema_ctis"
    date_basis = "public_application_submission_or_decision_date"
    coverage = "Public CTIS records only; historical public versions unavailable. Search limited to 10,000 results; incomplete ranges are partial."

    def page(self, client, start, end, mode, state, size):
        from ema_ctis_actions import SEARCH_URL, RETRIEVE_URL, PORTAL_URL, classify_decision
        size = min(size, 100)
        number = state.get("page", 1)
        if (number - 1) * size >= 10000:
            raise SourceIssue("ctis_public_search_limit", True)
        response = client.json(SEARCH_URL, body={"pagination": {"page": number, "size": size},
            "sort": {"property": "decisionDate", "direction": "DESC"}, "searchCriteria": {}})
        rows, pagination = response.get("data"), response.get("pagination", {})
        require(isinstance(rows, list) and isinstance(pagination.get("nextPage"), bool), "invalid_ctis_pagination")
        total, count = validate_total({"totalCount": pagination.get("totalRecords")}, rows, state)
        dates = [day(row.get("decisionDateOverall")) for row in rows]
        require(all(dates), "missing_ctis_decision_date")
        require(dates == sorted(dates, reverse=True) and (not state.get("last_date") or not dates or max(dates) <= state["last_date"]), "ctis_sort_changed")
        output = []
        for summary, decision_date in zip(rows, dates):
            if decision_date < start.isoformat():
                continue
            identity = summary.get("ctNumber")
            require(identity and re.fullmatch(r"\d{4}-\d{6}-\d{2}-\d{2}", identity), "missing_ctis_identifier")
            detail = client.json(RETRIEVE_URL.format(ct_number=identity))
            apps = detail.get("authorizedApplication", {}).get("applicationInfo")
            require(isinstance(apps, list), "missing_ctis_applications")
            events = []
            url = PORTAL_URL.format(ct_number=identity)
            for app in apps:
                require(app.get("id") is not None, "missing_ctis_application_id")
                submitted = day(app.get("submissionDate"))
                if app.get("submissionDate"):
                    require(submitted, "invalid_ctis_submission_date")
                if in_window(submitted, start, end):
                    events.append(item(f"{identity}:application:{app['id']}:submitted", "regulatory_event",
                        {"action": "SUBMITTED", "application": app, "trial_summary": summary}, url, self.date_basis,
                        submitted, [summary.get("sponsor")], summary.get("ctTitle"), identity_basis="documented_composite", coverage_note=self.coverage))
                decisions = app.get("decisions", [])
                require(isinstance(decisions, list), "invalid_ctis_decisions")
                for decision in decisions:
                    event = day(decision.get("decisionDate"))
                    if decision.get("decisionDate"):
                        require(event, "invalid_ctis_decision_date")
                    if in_window(event, start, end):
                        require(decision.get("id") is not None, "missing_ctis_decision_id")
                        events.append(item(f"{identity}:decision:{decision['id']}", "regulatory_event",
                            {"action": classify_decision(decision.get("decision")) or "OTHER_DECISION",
                             "application_id": app["id"], "decision": decision, "trial_summary": summary}, url,
                            self.date_basis, event, [summary.get("sponsor")], summary.get("ctTitle"),
                            identity_basis="documented_composite", coverage_note=self.coverage))
            if events:
                output.append(item(identity, "clinical_trial", {"detail": detail, "summary": summary}, url,
                    "current_trial_with_event_in_window", None, [summary.get("sponsor")], summary.get("ctTitle"),
                    updated=day(summary.get("lastUpdated")), coverage_note=self.coverage))
                output.extend(events)
        boundary = bool(dates) and min(dates) < start.isoformat()
        if not pagination["nextPage"]:
            require(count == total, "ctis_count_mismatch")
        return output, {"page": number + 1, "count": count, "total": total,
                        "last_date": min(dates) if dates else None}, boundary or not pagination["nextPage"]


def field_names(fields):
    return list(dict.fromkeys(f["value"] for f in fields
        if any(word in f["label"].lower() for word in ("sponsor", "申办", "申请人名称", "注册申请单位"))
        and not any(word in f["label"].lower() for word in ("contact", "联系人", "电话", "email", "地址"))
        and not f["value"].isdigit()))


def render_documents(client, urls):
    """Bounded detail concurrency, preserving listing order and page atomicity."""
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=3) as pool:
        return list(pool.map(client.rendered, urls))


class ChiCTR:
    source = "chictr"
    date_basis = "registration_date"
    coverage = "Registration discovery by exhaustive year-page traversal; listing dates need not be sorted. Reconcile old registration windows to observe revisions; no historical versions are reconstructed."

    def page(self, client, start, end, mode, state, size):
        import chictr_studies as legacy
        year, number = state.get("year", start.year), state.get("page", 1)
        url = legacy.SEARCH_URL + "?" + urlencode({"page": number, "btngo": "btn", "createyear": year})
        document = client.rendered(url)
        rows, total_pages = legacy.parse_search_page(document)
        # The current rendered page exposes a record total instead of a
        # textual page count. Preserve first-page size across the last page.
        count_match = re.search(r'id=["\']data-total["\'][^>]*>\s*(\d+)', document)
        total_records = int(count_match.group(1)) if count_match else None
        registry_size = state.get("registry_page_size", len(rows))
        if total_pages is None and total_records is not None and registry_size:
            total_pages = (total_records + registry_size - 1) // registry_size
        if total_records == 0 and not rows and number == 1:
            return [], {"year": year + 1, "page": 1}, year == (end - timedelta(days=1)).year
        require(rows and total_pages is not None and number <= total_pages, "chictr_pagination_unverified")
        if state.get("total_pages") is not None:
            require(total_pages == state["total_pages"], "source_changed_during_pagination")
        if state.get("total_records") is not None:
            require(total_records == state["total_records"], "source_changed_during_pagination")
        page_hash = digest(rows)
        require(page_hash != state.get("page_hash"), "pagination_cycle")
        dates = [row["registrationDate"] for row in rows]
        output = []
        selected = [row for row in rows if in_window(row["registrationDate"], start, end)]
        details = iter(render_documents(client, [row["detailUrl"] for row in selected]))
        for row in rows:
            require(row["registrationDate"][:4] == str(year), "chictr_year_filter_ignored")
            if in_window(row["registrationDate"], start, end):
                fields = legacy.parse_detail_page(next(details))
                require(fields, "empty_trial_detail")
                require(any(row["registrationNumber"] in field["value"] for field in fields
                            if "registration number" in field["label"].lower() or "注册号" in field["label"]),
                        "chictr_detail_identity_mismatch")
                output.append(item(row["registrationNumber"], "clinical_trial", {**row, "fields": fields}, row["detailUrl"],
                    self.date_basis, row["registrationDate"], field_names(fields), row["title"], coverage_note=self.coverage))
        # Registry ordering can differ from registration date. A page with
        # older dates cannot prove that later pages contain no matching rows.
        year_done = number == total_pages
        exhausted = year_done and year == (end - timedelta(days=1)).year
        if year_done:
            following = {"year": year + 1, "page": 1}
        else:
            following = {"year": year, "page": number + 1, "total_pages": total_pages, "page_hash": page_hash,
                         "registry_page_size": registry_size, "total_records": total_records, "last_date": min(dates)}
        return output, following, exhausted


def nmpa_summaries(document):
    """Current listing has no date column: dates must come from trial details."""
    from chictr_studies import TableParser
    from nmpa_drug_trials import BASE_URL, CTR_NUMBER
    parser = TableParser()
    parser.feed(document)
    output = []
    for row in parser.rows:
        text = " | ".join(cell["text"] for cell in row)
        match = CTR_NUMBER.search(text)
        if not match:
            continue
        links = [link["href"] for cell in row for link in cell["links"]
                 if link.get("href") and not link["href"].lower().startswith("javascript:")]
        number = match.group(0).upper()
        detail = next((link for link in links if "searchlistdetail.dhtml" in link),
                      "clinicaltrials.searchlistdetail.dhtml?" + urlencode({"reg_no": number}))
        output.append({"registrationNumber": number, "detailUrl": urljoin(BASE_URL, detail),
                       "title": row[-1]["text"], "searchResultColumns": [cell["text"] for cell in row]})
    return output


class NMPA:
    source = "nmpa_trials"
    date_basis = "first_public_disclosure_date"
    coverage = "First-disclosure discovery, not an update feed. Dates come from detail records; full listing pagination is required. Reconcile older windows for changes."

    def __init__(self, registry_ids=()):
        self.registry_ids = tuple(sorted(set(registry_ids)))
        if self.registry_ids:
            self.coverage = "Explicit registration-ID verification cohort only; not complete date-window or registry coverage. Current trial details, filtered by first public disclosure date."

    def page(self, client, start, end, mode, state, size):
        import nmpa_drug_trials as legacy
        number = state.get("page", 1)
        if self.registry_ids:
            require(1 <= number <= len(self.registry_ids), "invalid_cohort_cursor")
            registration = self.registry_ids[number - 1]
            url = legacy.BASE_URL + "clinicaltrials.searchlistdetail.dhtml?" + urlencode({"reg_no": registration})
            rows = [{"registrationNumber": registration, "detailUrl": url, "title": None}]
            total_pages = len(self.registry_ids)
        else:
            return self.list_page(client, start, end, mode, state, size)
        output = self.details(client, rows, start, end)
        return output, {"page": number + 1}, number == total_pages

    def list_page(self, client, start, end, mode, state, size):
        import nmpa_drug_trials as legacy
        number = state.get("page", 1)
        document = client.rendered(legacy.SEARCH_URL + "?" + urlencode({"keywords": "", "currentpage": number}))
        rows = nmpa_summaries(document)
        require(rows, "nmpa_empty_or_unparseable_page")
        page_hash = digest(rows)
        require(page_hash != state.get("page_hash"), "pagination_cycle")
        visible = unescape(re.sub(r"<[^>]+>", " ", document))
        current = re.search(r"当前第\s*(\d+)\s*页", visible)
        require(current is not None and int(current.group(1)) == number, "nmpa_page_parameter_ignored")
        pages_match = re.search(r"共\s*(\d+)\s*页", visible)
        require(pages_match is not None, "nmpa_pagination_unverified")
        total_pages = int(pages_match.group(1))
        require(number <= total_pages and ("total_pages" not in state or state["total_pages"] == total_pages), "nmpa_pagination_changed")
        total_match = re.search(r"共\s*(\d+)\s*条记录", visible)
        require(total_match is not None, "nmpa_record_total_missing")
        total_records = int(total_match.group(1))
        require("total_records" not in state or total_records == state["total_records"], "source_changed_during_pagination")
        output = self.details(client, rows, start, end)
        return output, {"page": number + 1, "total_pages": total_pages, "page_hash": page_hash,
                        "total_records": total_records}, number == total_pages

    def details(self, client, rows, start, end):
        import nmpa_drug_trials as legacy
        output = []
        for row, document in zip(rows, render_documents(client, [row["detailUrl"] for row in rows])):
            fields = legacy.parse_detail_page(document)
            require(fields, "empty_trial_detail")
            text = unescape(re.sub(r"<[^>]+>", " ", document))
            require(any(field["label"] == "登记号" and
                        re.search(r"\b" + re.escape(row["registrationNumber"]) + r"\b", field["value"])
                        for field in fields), "nmpa_detail_identity_mismatch")
            match = re.search(r"首次公示(?:信息)?日期[：:\s]*(20\d{2}[-/.]\d{1,2}[-/.]\d{1,2})", text)
            parsed = legacy.parse_registry_date(match.group(1)) if match else None
            require(parsed is not None, "nmpa_disclosure_date_missing")
            row["firstPublicDisclosureDate"] = parsed.isoformat()
            if not row.get("title"):
                row["title"] = next((field["value"] for field in fields if field["label"] in ("试验通俗题目", "试验专业题目")), None)
            if in_window(row["firstPublicDisclosureDate"], start, end):
                output.append(item(row["registrationNumber"], "clinical_trial", {**row, "fields": fields}, row["detailUrl"],
                    self.date_basis, row["firstPublicDisclosureDate"], field_names(fields), row["title"], coverage_note=self.coverage))
        return output


class UK:
    source = "uk_mhra"
    date_basis = "current_snapshot_with_optional_record_date"
    coverage = "Current public lists, including undated rows. Requested windows do not recreate historical snapshots; no added/removed events are inferred."

    def page(self, client, start, end, mode, state, size):
        import uk_mhra_designations as legacy
        from openpyxl import load_workbook
        output = []
        sources = [("ilap", legacy.ILAP_PAGE, ".xlsx", "Innovation Passport"),
                   ("orphan", legacy.ORPHAN_PAGE, ".csv", "Orphan Register")]
        stage = state.get("stage", 0)
        category, page, suffix, label = sources[stage]
        parser = legacy.Links()
        parser.feed(client.get(page).decode("utf-8"))
        matches = [urljoin(page, href) for href, text in parser.links
                   if href.lower().split("?")[0].endswith(suffix) and label.lower() in text.lower()]
        require(matches, "mhra_attachment_missing")
        url = matches[-1]
        raw = client.get(url)
        if suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
            header_at = next((i for i, row in enumerate(rows) if sum(v is not None for v in row) >= 2), None)
            require(header_at is not None, "mhra_header_missing")
            headers = [str(v).strip() if v is not None else f"column_{i+1}" for i, v in enumerate(rows[header_at])]
            records = [{headers[i]: legacy.clean(v) for i, v in enumerate(row) if i < len(headers) and v not in (None, "")}
                       for row in rows[header_at+1:] if any(v not in (None, "") for v in row)]
        else:
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1252")
            lines = text.splitlines()
            header_at = next((i for i, line in enumerate(lines) if line.lstrip().lower().startswith("product name")), None)
            require(header_at is not None, "mhra_header_missing")
            records = [{str(k).strip(): legacy.clean(v) for k, v in row.items() if k and v not in (None, "")}
                       for row in csv.DictReader(io.StringIO("\n".join(lines[header_at:])))]
        require(records, "mhra_empty_snapshot")
        for row in records:
            lower = {key.lower().strip(" *"): str(value) for key, value in row.items()}
            identity_parts = [lower[key] for key in ("designation number", "product name", "company name") if lower.get(key)]
            require(identity_parts, "mhra_identity_columns_missing")
            identity = category + ":" + digest([x.casefold() for x in identity_parts])
            event = legacy.record_date(row)
            # Daily preserves the current list, including undated rows.
            # Historical windows can only select known dates in that list.
            if mode == "historical":
                if event is None:
                    raise SourceIssue("historical_designation_date_unavailable", True)
                if not start <= event < end:
                    continue
            output.append(item(identity, "regulatory_snapshot", {"category": category, "record": row}, url,
                self.date_basis, str(event) if event else None,
                [value for key, value in lower.items() if any(x in key for x in ("company", "holder", "applicant"))],
                lower.get("product name"), identity_basis="documented_composite", coverage_note=self.coverage))
        return output, {"stage": stage + 1}, stage == 1


class Japan:
    source = "japan_pmda"
    date_basis = "current_sakigake_list_designation_date"
    coverage = "Current published SAKIGAKE PDF; not historical snapshots or complete PMDA regulatory coverage."

    def page(self, client, start, end, mode, state, size):
        import japan_pmda_designations as legacy
        records = legacy.parse_pdf(client.get(legacy.SOURCE_URL))
        output = [item("sakigake:" + row["designationDate"] + ":" + row["listNumber"], "regulatory_snapshot", row, legacy.SOURCE_URL,
            self.date_basis, row["designationDate"], (), row["recordText"], identity_basis="documented_composite",
            coverage_note=self.coverage) for row in records
            if mode != "historical" or in_window(row["designationDate"], start, end)]
        return output, {}, True


ADAPTERS = {"clinicaltrials": ClinicalTrials, "fda": FDA, "ema": EMA, "chictr": ChiCTR,
            "nmpa": NMPA, "uk": UK, "japan": Japan}
